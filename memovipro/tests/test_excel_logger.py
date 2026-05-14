from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from core.excel_logger import ExcelLogger, Incidencia, HEADERS


def test_logger_crea_y_acumula(tmp_path):
    path = tmp_path / "incidencias.xlsx"
    logger = ExcelLogger(path)
    assert path.exists()

    logger.append(Incidencia(
        dni="12345678A", macro="m1", paso_idx=3, paso_tipo="click_control",
        tipo_error="POPUP", titulo_popup="Error", texto_popup="DNI no encontrado",
    ))
    logger.append_ok("99999999B", "m1", detalle="ok")

    wb = load_workbook(path)
    ws = wb["Incidencias"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == tuple(HEADERS)
    assert rows[1][1] == "12345678A"
    assert rows[1][HEADERS.index("Estado")] == "ERROR"
    assert rows[2][1] == "99999999B"
    assert rows[2][HEADERS.index("Estado")] == "OK"


def test_count_by_dni(tmp_path):
    path = tmp_path / "i.xlsx"
    logger = ExcelLogger(path)
    for _ in range(3):
        logger.append(Incidencia(dni="A", macro="m", paso_idx=0, paso_tipo="x", tipo_error="POPUP"))
    logger.append(Incidencia(dni="B", macro="m", paso_idx=0, paso_tipo="x", tipo_error="POPUP"))
    counts = logger.count_by_dni()
    assert counts == {"A": 3, "B": 1}
