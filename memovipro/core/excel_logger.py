from __future__ import annotations

import csv
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from threading import Lock

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from loguru import logger
except Exception:  # pragma: no cover
    class _NullLogger:
        def warning(self, *a, **kw): pass
        def error(self, *a, **kw): pass
    logger = _NullLogger()


HEADERS = [
    "Fecha/Hora",
    "DNI",
    "Macro",
    "Paso #",
    "Tipo paso",
    "Tipo error",
    "Título popup",
    "Texto popup",
    "Screenshot",
    "Estado",
    "Detalle",
]

HEADER_FILL = PatternFill("solid", fgColor="2C3E50")
HEADER_FONT = Font(bold=True, color="FFFFFF")
OK_FILL = PatternFill("solid", fgColor="D5F5E3")
ERR_FILL = PatternFill("solid", fgColor="FADBD8")


@dataclass
class Incidencia:
    dni: str
    macro: str
    paso_idx: int
    paso_tipo: str
    tipo_error: str
    titulo_popup: str = ""
    texto_popup: str = ""
    screenshot_path: str = ""
    estado: str = "ERROR"
    detalle: str = ""
    timestamp: datetime | None = None

    def row(self) -> list:
        ts = self.timestamp or datetime.now()
        return [
            ts.strftime("%Y-%m-%d %H:%M:%S"),
            self.dni,
            self.macro,
            self.paso_idx,
            self.paso_tipo,
            self.tipo_error,
            self.titulo_popup,
            self.texto_popup,
            self.screenshot_path,
            self.estado,
            self.detalle,
        ]


class ExcelLogger:
    """Acumula incidencias en un Excel diario. Thread-safe."""

    def __init__(self, path: str | Path):
        self.path = Path(path)
        self._lock = Lock()
        self._ensure_workbook()

    def _ensure_workbook(self) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if not self.path.exists():
            wb = Workbook()
            ws = wb.active
            ws.title = "Incidencias"
            ws.append(HEADERS)
            for col_idx, _ in enumerate(HEADERS, start=1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")
            self._autosize(ws)
            ws.freeze_panes = "A2"
            wb.save(self.path)

    @staticmethod
    def _autosize(ws) -> None:
        widths = [20, 14, 24, 8, 22, 18, 30, 60, 50, 12, 40]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # Cuántas veces reintentar el guardado del Excel antes de rendirse y
    # volcar la fila al CSV de emergencia. Cubre el caso típico: el usuario
    # tiene el fichero abierto en Excel (bloqueo de escritura de Windows).
    _SAVE_RETRIES = 4

    def append(self, inc: Incidencia) -> int:
        """Añade una incidencia. Devuelve el número de fila escrita, o -1 si
        no se pudo escribir en el Excel (en cuyo caso la fila se vuelca a un
        CSV de emergencia y la ejecución del lote continúa).

        NUNCA propaga: perder el registro de UNA incidencia no debe abortar
        el procesamiento del resto de DNIs. El caso frecuente es tener el
        Excel abierto en Excel (bloqueo de escritura) — antes eso mataba la
        pasada entera.

        Protegido con un lock intra-proceso (threading) Y entre procesos
        (file_lock) para que una tarea programada y la GUI no corrompan el
        Excel escribiendo a la vez. El guardado es atómico (fichero temporal
        + reemplazo) para que un corte a mitad no deje el .xlsx corrupto.
        """
        from .file_lock import file_lock
        with self._lock:
            ultimo_err: Exception | None = None
            for intento in range(self._SAVE_RETRIES):
                try:
                    with file_lock(self.path):
                        return self._escribir_incidencia(inc)
                except PermissionError as exc:
                    ultimo_err = exc
                    # Fichero bloqueado (abierto en Excel): esperar y reintentar.
                    time.sleep(min(0.5 * (2 ** intento), 4.0))
                except Exception as exc:
                    ultimo_err = exc
                    logger.warning("Fallo al escribir incidencia (intento {}): {}", intento + 1, exc)
                    time.sleep(min(0.3 * (2 ** intento), 2.0))
            # Agotados los reintentos: no perder el dato ni abortar el lote.
            self._volcar_a_csv_emergencia(inc, ultimo_err)
            return -1

    def _escribir_incidencia(self, inc: Incidencia) -> int:
        """Carga el workbook, añade la fila y guarda de forma atómica."""
        wb = load_workbook(self.path)
        ws = wb["Incidencias"]
        row = inc.row()
        ws.append(row)
        row_idx = ws.max_row
        fill = OK_FILL if inc.estado.upper() == "OK" else ERR_FILL
        for col in range(1, len(HEADERS) + 1):
            ws.cell(row=row_idx, column=col).fill = fill
        if inc.screenshot_path:
            p = Path(inc.screenshot_path)
            if p.exists():
                cell = ws.cell(row=row_idx, column=HEADERS.index("Screenshot") + 1)
                cell.hyperlink = str(p.resolve())
                cell.value = p.name
                cell.font = Font(color="1A5276", underline="single")
        self._guardar_atomico(wb)
        return row_idx

    def _guardar_atomico(self, wb) -> None:
        """Guarda a un .tmp y lo reemplaza. Evita dejar el .xlsx corrupto
        si el proceso muere a mitad del save (todos los append posteriores
        fallarían para siempre porque load_workbook no puede abrirlo).

        Si el `replace` falla (destino bloqueado por Excel), limpiamos el
        .tmp para no dejar residuos huérfanos en data/ y re-lanzamos para
        que el bucle de reintentos de `append` lo gestione."""
        tmp = self.path.with_suffix(self.path.suffix + ".tmp")
        try:
            wb.save(tmp)
            tmp.replace(self.path)
        except Exception:
            try:
                tmp.unlink(missing_ok=True)
            except Exception:
                pass
            raise

    def _volcar_a_csv_emergencia(self, inc: Incidencia, err: Exception | None) -> None:
        """Escribe la fila en incidencias_YYYYMMDD.pendientes.csv cuando el
        Excel no se pudo escribir. Best-effort: si esto también falla, solo
        lo logueamos (jamás propagamos)."""
        csv_path = self.path.with_suffix(".pendientes.csv")
        try:
            nuevo = not csv_path.exists()
            with csv_path.open("a", encoding="utf-8-sig", newline="") as f:
                w = csv.writer(f)
                if nuevo:
                    w.writerow(HEADERS)
                w.writerow(inc.row())
            logger.error(
                "No se pudo escribir en {} ({}). Incidencia volcada a {}",
                self.path.name, err, csv_path.name,
            )
        except Exception as exc:  # pragma: no cover - fallo del último recurso
            logger.error("No se pudo ni volcar la incidencia al CSV de emergencia: {}", exc)

    def append_ok(self, dni: str, macro: str, detalle: str = "") -> int:
        return self.append(
            Incidencia(
                dni=dni,
                macro=macro,
                paso_idx=-1,
                paso_tipo="",
                tipo_error="",
                estado="OK",
                detalle=detalle,
            )
        )

    def count_by_dni(self) -> dict[str, int]:
        with self._lock:
            wb = load_workbook(self.path, read_only=True)
            try:
                ws = wb["Incidencias"]
                counts: dict[str, int] = {}
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or not row[1]:
                        continue
                    counts[row[1]] = counts.get(row[1], 0) + 1
                return counts
            finally:
                # En Windows un workbook read_only sin cerrar mantiene el
                # fichero bloqueado.
                wb.close()


def default_log_path(carpeta: str | Path, patron: str = "incidencias_{YYYYMMDD}.xlsx") -> Path:
    name = patron.replace("{YYYYMMDD}", datetime.now().strftime("%Y%m%d"))
    return Path(carpeta) / name
