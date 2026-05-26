"""Agregación de las incidencias históricas para el dashboard.

Lee todos los `incidencias_YYYYMMDD.xlsx` de la carpeta data/ y produce
estadísticas: por macro (OK/KO/tasa de éxito) y los errores más
frecuentes. Pensado para que el usuario vea de un vistazo qué macros son
frágiles antes de que le fallen en producción.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook


_FECHA_RE = re.compile(r"incidencias_(\d{8})\.xlsx$", re.IGNORECASE)


@dataclass
class MacroStats:
    macro: str
    ok: int = 0
    ko: int = 0

    @property
    def total(self) -> int:
        return self.ok + self.ko

    @property
    def tasa_exito(self) -> float:
        return (self.ok / self.total * 100.0) if self.total else 0.0


@dataclass
class DashboardStats:
    desde: str = ""
    hasta: str = ""
    archivos_leidos: int = 0
    por_macro: list[MacroStats] = field(default_factory=list)
    top_errores: list[tuple[str, int]] = field(default_factory=list)
    total_ok: int = 0
    total_ko: int = 0

    @property
    def total(self) -> int:
        return self.total_ok + self.total_ko


def _fecha_de_archivo(path: Path) -> datetime | None:
    m = _FECHA_RE.search(path.name)
    if not m:
        return None
    try:
        return datetime.strptime(m.group(1), "%Y%m%d")
    except ValueError:
        return None


def agregar_incidencias(data_dir: str | Path, dias: int = 30) -> DashboardStats:
    """Agrega los incidencias_*.xlsx de los últimos `dias` días.

    Columnas esperadas (de excel_logger.HEADERS):
      Fecha/Hora, DNI, Macro, Paso #, Tipo paso, Tipo error,
      Título popup, Texto popup, Screenshot, Estado, Detalle
    """
    data_dir = Path(data_dir)
    stats = DashboardStats()
    if not data_dir.exists():
        return stats

    limite = datetime.now() - timedelta(days=dias)
    archivos = []
    for p in sorted(data_dir.glob("incidencias_*.xlsx")):
        f = _fecha_de_archivo(p)
        if f is None or f < limite:
            continue
        archivos.append((f, p))

    macros: dict[str, MacroStats] = {}
    errores: dict[str, int] = {}
    fechas: list[datetime] = []

    for fecha, path in archivos:
        try:
            wb = load_workbook(path, read_only=True)
            ws = wb["Incidencias"]
        except Exception:
            continue
        stats.archivos_leidos += 1
        fechas.append(fecha)
        # Índices de columna (0-based) según HEADERS
        IDX_MACRO, IDX_ERR, IDX_TITULO, IDX_ESTADO, IDX_DETALLE = 2, 5, 6, 9, 10
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not any(row):
                continue
            macro = (row[IDX_MACRO] if len(row) > IDX_MACRO else "") or "(sin macro)"
            estado = (row[IDX_ESTADO] if len(row) > IDX_ESTADO else "") or ""
            estado = str(estado).strip().upper()
            ms = macros.setdefault(macro, MacroStats(macro=macro))
            if estado == "OK":
                ms.ok += 1
                stats.total_ok += 1
            else:
                ms.ko += 1
                stats.total_ko += 1
                # Agregar el motivo del error (título popup o detalle)
                titulo = (row[IDX_TITULO] if len(row) > IDX_TITULO else "") or ""
                detalle = (row[IDX_DETALLE] if len(row) > IDX_DETALLE else "") or ""
                motivo = str(titulo).strip() or str(detalle).strip() or "(sin detalle)"
                motivo = motivo[:120]
                errores[motivo] = errores.get(motivo, 0) + 1

    stats.por_macro = sorted(
        macros.values(), key=lambda m: (m.tasa_exito, -m.total)
    )  # frágiles (peor tasa) primero
    stats.top_errores = sorted(errores.items(), key=lambda kv: -kv[1])[:15]
    if fechas:
        stats.desde = min(fechas).strftime("%Y-%m-%d")
        stats.hasta = max(fechas).strftime("%Y-%m-%d")
    return stats
