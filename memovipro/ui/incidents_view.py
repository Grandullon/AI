from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QColor, QDesktopServices
from PyQt6.QtCore import QUrl
from PyQt6.QtWidgets import (
    QComboBox,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from core.excel_logger import HEADERS


class IncidentsView(QWidget):
    def __init__(self, data_dir: Path):
        super().__init__()
        self.data_dir = data_dir

        layout = QVBoxLayout(self)

        top = QHBoxLayout()
        self.file_combo = QComboBox()
        refresh = QPushButton("↻ Refrescar")
        refresh.clicked.connect(self.refresh)
        abrir = QPushButton("Abrir Excel")
        abrir.clicked.connect(self._open_excel)
        top.addWidget(QLabel("Archivo:"))
        top.addWidget(self.file_combo, 1)
        top.addWidget(refresh)
        top.addWidget(abrir)
        layout.addLayout(top)

        self.tabla = QTableWidget(0, len(HEADERS))
        self.tabla.setHorizontalHeaderLabels(HEADERS)
        self.tabla.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        self.tabla.setAlternatingRowColors(True)
        self.tabla.verticalHeader().setVisible(False)
        self.tabla.cellDoubleClicked.connect(self._on_double_click)
        layout.addWidget(self.tabla, 1)

        self.summary = QLabel("")
        layout.addWidget(self.summary)

        self.file_combo.currentIndexChanged.connect(self._load_selected)
        self.refresh()

    def refresh(self):
        self.file_combo.clear()
        if not self.data_dir.exists():
            return
        files = sorted(self.data_dir.glob("incidencias_*.xlsx"), reverse=True)
        for f in files:
            self.file_combo.addItem(f.name, str(f))
        self._load_selected()

    def _load_selected(self):
        path = self.file_combo.currentData()
        self.tabla.setRowCount(0)
        if not path or not Path(path).exists():
            self.summary.setText("Sin archivos de incidencias")
            return
        wb = load_workbook(path, read_only=True)
        ws = wb["Incidencias"]
        ok = 0
        ko = 0
        dnis_fallidos = set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not any(row):
                continue
            r = self.tabla.rowCount()
            self.tabla.insertRow(r)
            for col, val in enumerate(row[: len(HEADERS)]):
                item = QTableWidgetItem("" if val is None else str(val))
                self.tabla.setItem(r, col, item)
            estado = (row[HEADERS.index("Estado")] or "").strip().upper() if row[HEADERS.index("Estado")] else ""
            color = QColor("#D5F5E3") if estado == "OK" else QColor("#FADBD8")
            for col in range(len(HEADERS)):
                self.tabla.item(r, col).setBackground(color)
            if estado == "OK":
                ok += 1
            else:
                ko += 1
                dni = row[HEADERS.index("DNI")]
                if dni:
                    dnis_fallidos.add(dni)
        self.summary.setText(
            f"Total filas: {ok + ko}  ·  OK: {ok}  ·  KO: {ko}  ·  DNIs distintos con fallo: {len(dnis_fallidos)}"
        )

    def _on_double_click(self, row: int, col: int):
        if HEADERS[col] != "Screenshot":
            return
        item = self.tabla.item(row, col)
        if not item:
            return
        path = item.text()
        if not path:
            return
        QDesktopServices.openUrl(QUrl.fromLocalFile(path))

    def _open_excel(self):
        path = self.file_combo.currentData()
        if path:
            QDesktopServices.openUrl(QUrl.fromLocalFile(path))
